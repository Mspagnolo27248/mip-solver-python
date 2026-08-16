"""Import the FY26 planning workbook into JSON seed files.

Produces two artifacts:

  reference.json  - master/reference data that changes rarely (products, capacities,
                    yields, run rates, blend BOM, targets, block specs)
  scenario.json   - one planning scenario: the feeds (inventory, orders, forecast)
                    plus the charge schedule that the planner entered

Everything the simulation engine needs comes from these two files, so the engine
never touches Excel and the same code path serves a database-backed web app later.
"""
from __future__ import annotations

import datetime as dt
import json
import os
import re
from collections import defaultdict
from typing import Any, Dict, List, Optional, Tuple

from . import layout as L
from .xlsx import (GAL_PER_BBL, as_date, formula, is_num, load, month_key,
                   norm_code, num, read_date_row, text)


class ImportWarnings:
    def __init__(self) -> None:
        self.items: List[str] = []

    def add(self, category: str, message: str) -> None:
        self.items.append("[{}] {}".format(category, message))

    def __len__(self) -> int:
        return len(self.items)


def _iso(d: dt.date) -> str:
    return d.isoformat()


# Rows whose value is sourced from feeds/schedule (as opposed to being derived
# from other rows in the same block, which the engine computes directly).
_CLASSIFIED_LABELS = {
    "Production In", "Receipts", "Sales", "Blends", "Forecast",
    "Production Out", "Production Out &DG DFO", "Out to Diesel", "Downgrade",
}


class WorkbookImporter:
    def __init__(self, path: str):
        self.path = path
        self.wf = load(path, data_only=False)   # formulas
        self.wv = load(path, data_only=True)    # cached values
        self.warn = ImportWarnings()
        self._date_cols_cache: Dict[str, List[Tuple[dt.date, int]]] = {}

    # ------------------------------------------------------------------ helpers
    def fv(self, sheet: str, row: int, col: int) -> Any:
        return self.wv[sheet].cell(row, col).value

    def ff(self, sheet: str, row: int, col: int) -> str:
        return formula(self.wf[sheet], row, col)

    def validate_layout(self) -> None:
        """Fail loudly if the anchors in layout.py no longer point where expected."""
        checks: List[Tuple[str, Any, str]] = [
            ("Charge Schedule crude row",
             text(self.fv(L.CS, L.CS_CRUDE_BBL_ROW, 2)), "Crude Oil Charge"),
            ("Production-Out loss row",
             text(self.fv(L.PO, L.PO_LOSS_ROW, 4)), "Loss"),
            ("Production-Out gallons grid",
             text(self.fv(L.PO, L.PO_GAL_DATE_ROW, 3)), "Code"),
            ("Open Orders header",
             text(self.fv(L.OPEN_ORDERS, 8, 1)), "Code"),
        ]
        for what, got, expect in checks:
            if got != expect:
                raise RuntimeError(
                    "Layout drift: {} expected {!r} but found {!r}. "
                    "Update layout.py before importing.".format(what, expect, got))

        for blk in L.CHARGE_BLOCKS:
            for r in range(blk.first_row, blk.last_row + 1):
                if self.fv(L.CS, r, 2) is None and self.fv(L.CS, r, 3) is None:
                    self.warn.add("layout",
                                  "charge block {} row {} is empty".format(blk.unit, r))

    # --------------------------------------------------------------- reference
    def products(self) -> Dict[str, Dict[str, str]]:
        ws = self.wv[L.PRODUCT_LIST]
        out: Dict[str, Dict[str, str]] = {}
        for r in range(L.PRODUCT_LIST_FIRST_ROW, ws.max_row + 1):
            code = norm_code(self.fv(L.PRODUCT_LIST, r, 1))
            full = text(self.fv(L.PRODUCT_LIST, r, 2))
            name = text(self.fv(L.PRODUCT_LIST, r, 4))
            if not code or not (full or name):
                continue
            out.setdefault(code, {"name": name or full, "full": full})
        return out

    def tank_capacity(self) -> Dict[str, float]:
        ws = self.wv[L.TANK_CAP]
        out: Dict[str, float] = {}
        for r in range(L.TANK_CAP_FIRST_ROW, ws.max_row + 1):
            code = norm_code(self.fv(L.TANK_CAP, r, L.TANK_CAP_CODE_COL))
            val = self.fv(L.TANK_CAP, r, L.TANK_CAP_VALUE_COL)
            if not code or not is_num(val):
                continue
            if code in out and out[code] != float(val):
                self.warn.add("tank_capacity",
                              "duplicate capacity rows for {} ({} vs {}); keeping first"
                              .format(code, out[code], val))
                continue
            out[code] = float(val)
        return out

    def inventory_targets(self) -> Dict[str, Dict[str, Optional[float]]]:
        out: Dict[str, Dict[str, Optional[float]]] = {}
        for r in range(L.IT_FIRST_ROW, L.IT_LAST_ROW + 1):
            code = norm_code(self.fv(L.INV_TARGETS, r, L.IT_CODE_COL))
            if not code:
                continue
            lcl = self.fv(L.INV_TARGETS, r, L.IT_LCL_COL)
            ucl = self.fv(L.INV_TARGETS, r, L.IT_UCL_COL)
            if not is_num(lcl) and not is_num(ucl):
                continue
            out[code] = {"lcl": float(lcl) if is_num(lcl) else None,
                         "ucl": float(ucl) if is_num(ucl) else None,
                         "name": text(self.fv(L.INV_TARGETS, r, L.IT_NAME_COL))}
        return out

    def _month_header(self, sheet: str, row: int, first_col: int, count: int = 12
                      ) -> List[str]:
        months = []
        for i in range(count):
            v = text(self.fv(sheet, row, first_col + i)).upper()[:3]
            months.append(v)
        return months

    def crude_yields(self) -> Tuple[Dict[str, Dict[str, float]], float]:
        months = self._month_header(L.CRUDE_YIELDS, L.CY_MONTH_HEADER_ROW,
                                    L.CY_FIRST_MONTH_COL)
        out: Dict[str, Dict[str, float]] = {}
        loss = 0.0
        for r in range(L.CY_FIRST_ROW, L.CY_LAST_ROW + 1):
            code = norm_code(self.fv(L.CRUDE_YIELDS, r, L.CY_CODE_COL))
            name = text(self.fv(L.CRUDE_YIELDS, r, L.CY_NAME_COL))
            vals = {}
            for i, m in enumerate(months):
                v = self.fv(L.CRUDE_YIELDS, r, L.CY_FIRST_MONTH_COL + i)
                if is_num(v):
                    vals[m] = float(v)
            if not vals:
                continue
            if code:
                out[code] = vals
            elif name.lower() == "loss":
                loss = max(vals.values())
        return out, loss

    def run_rates(self) -> Dict[str, Dict[str, float]]:
        months = self._month_header(L.RUN_RATES, L.RR_MONTH_HEADER_ROW,
                                    L.RR_FIRST_MONTH_COL)
        out: Dict[str, Dict[str, float]] = {}
        for r in range(L.RR_FIRST_ROW, L.RR_LAST_ROW + 1):
            code = norm_code(self.fv(L.RUN_RATES, r, L.RR_CODE_COL))
            if not code:
                continue
            vals = {}
            for i, m in enumerate(months):
                v = self.fv(L.RUN_RATES, r, L.RR_FIRST_MONTH_COL + i)
                if is_num(v):
                    vals[m] = float(v)
            if vals:
                out.setdefault(code, vals)
        return out

    def other_rates(self) -> Dict[str, Dict[str, float]]:
        months = self._month_header(L.RUN_RATES, L.RR_MONTH_HEADER_ROW,
                                    L.RR_FIRST_MONTH_COL)
        res = {}
        for key, row in (("k30_to_diesel", L.RR_K30_TO_DIESEL_ROW),
                         ("dewaxed_lln", L.RR_DEWAXED_LLN_ROW)):
            vals = {}
            for i, m in enumerate(months):
                v = self.fv(L.RUN_RATES, row, L.RR_FIRST_MONTH_COL + i)
                if is_num(v):
                    vals[m] = float(v)
            res[key] = vals
        return res

    def crude_plan(self) -> Dict[str, Dict[str, float]]:
        charge: Dict[str, float] = {}
        first, last = L.CR_MONTH_ROWS
        for r in range(first, last + 1):
            m = text(self.fv(L.CRUDE_RATES, r, 2)).upper()[:3]
            v = self.fv(L.CRUDE_RATES, r, 3)
            if m and is_num(v):
                charge[m] = float(v)
        receipts: Dict[str, float] = {}
        for i in range(12):
            c = L.CR_RECEIPT_FIRST_COL + i
            m = text(self.fv(L.CRUDE_RATES, L.CR_RECEIPT_HDR_ROW, c)).upper()[:3]
            v = self.fv(L.CRUDE_RATES, L.CR_RECEIPT_ROW, c)
            if m and is_num(v):
                receipts[m] = float(v)
        return {"charge_bbl_per_day": charge, "receipts_bbl_per_day": receipts}

    def blend_bom(self) -> List[Dict[str, Any]]:
        out = []
        for r in range(L.BLENDS_BOM_FIRST_ROW, L.BLENDS_BOM_LAST_ROW + 1):
            frac = self.fv(L.BLENDS, r, L.BLENDS_FRACTION_COL)
            blend = norm_code(self.fv(L.BLENDS, r, L.BLENDS_BLEND_COL))
            comp = norm_code(self.fv(L.BLENDS, r, L.BLENDS_COMPONENT_COL))
            if not is_num(frac) or not blend or not comp:
                continue
            out.append({"blend": blend, "component": comp, "fraction": float(frac)})
        totals: Dict[str, float] = defaultdict(float)
        for row in out:
            totals[row["blend"]] += row["fraction"]
        for blend, tot in sorted(totals.items()):
            if abs(tot - 1.0) > 0.01:
                self.warn.add("blend_bom",
                              "BOM for blend {} sums to {:.4f}, not 1.0".format(blend, tot))
        return out

    # ----------------------------------------------------------- charge schedule
    def charge_dates(self) -> List[Tuple[int, dt.date]]:
        return read_date_row(self.wv[L.CS], L.CS_DATE_ROW_BBL, L.CS_FIRST_COL)

    def charge_lines(self) -> List[Dict[str, Any]]:
        """Every charge/transfer entry row, in sheet order (order matters: the
        stream sheets resolve 'Production Out' with a first-match VLOOKUP)."""
        lines = []
        for blk in L.CHARGE_BLOCKS:
            for i, r in enumerate(range(blk.first_row, blk.last_row + 1)):
                code = norm_code(self.fv(L.CS, r, 2))
                name = text(self.fv(L.CS, r, 3))
                if not code:
                    continue
                lines.append({
                    "key": "{}#{}".format(blk.unit, r),
                    "unit": blk.unit,
                    "row": r,
                    "code": code,
                    "name": name,
                    "index_in_block": i,
                    "in_charge_lookup_range":
                        L.CHARGE_LOOKUP_FIRST_ROW <= r <= L.CHARGE_LOOKUP_LAST_ROW,
                })
        return lines

    def charge_schedule(self, lines: List[Dict[str, Any]],
                        dates: List[Tuple[int, dt.date]]) -> Dict[str, Any]:
        crude_bbl: Dict[str, float] = {}
        crude_mode: Dict[str, str] = {}
        flag_dates = {d: c for c, d in
                      read_date_row(self.wv[L.CS], L.CS_DATE_ROW_FLAG, L.CS_FIRST_COL)}
        for c, d in dates:
            v = self.fv(L.CS, L.CS_CRUDE_BBL_ROW, c)
            if is_num(v) and v:
                crude_bbl[_iso(d)] = float(v)
            fc = flag_dates.get(d)
            if fc:
                mv = text(self.fv(L.CS, L.CS_CRUDE_MODE_ROW, fc)).upper()
                if mv in ("R", "L"):
                    crude_mode[_iso(d)] = mv

        entries: Dict[str, Dict[str, float]] = {}
        for line in lines:
            row = line["row"]
            per_date: Dict[str, float] = {}
            for c, d in dates:
                v = self.fv(L.CS, row, c)
                if is_num(v) and v:
                    per_date[_iso(d)] = float(v)
            if per_date:
                entries[line["key"]] = per_date
        return {"crude_bbl_per_day": crude_bbl,
                "crude_mode": crude_mode,
                "lines": entries}

    # ------------------------------------------------------------- yield rules
    def yield_rules(self, lines: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
        """Build the production rule table (BUSINESS-LOGIC section 3.2)."""
        rules: List[Dict[str, Any]] = []
        cy, loss = self.crude_yields()

        # --- crude unit side streams
        first, last = L.PO_CRUDE_ROWS
        for r in range(first, last + 1):
            code = norm_code(self.fv(L.PO, r, 3))
            if not code:
                continue
            f = self.ff(L.PO, r, L.PO_FIRST_COL)
            mode = None
            m = re.search(r'IF\(\s*E\d+\s*=\s*"([RL])"\s*,\s*0', f)
            if m:
                # `IF(mode="L",0,...)` means the stream is produced only in the OTHER mode
                mode = "R" if m.group(1) == "L" else "L"
            if code not in cy:
                self.warn.add("yields",
                              "crude row {} product {} has no CRUDE YIELDS entry".format(r, code))
                continue
            rules.append({"kind": "crude", "po_row": r, "out": code,
                          "monthly_yield": cy[code], "mode": mode})
        rules.append({"kind": "crude_loss", "po_row": L.PO_LOSS_ROW, "out": None,
                      "fraction": loss if loss < 1 else 1.0 - loss})

        # --- process units driven by charge lines
        by_unit_row: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
        for ln in lines:
            by_unit_row[ln["unit"]].append(ln)

        for sec in L.PO_UNIT_SECTIONS:
            for r in range(sec.first_row, sec.last_row + 1):
                out_code = norm_code(self.fv(L.PO, r, 3))
                charge_code = norm_code(self.fv(L.PO, r, 2))
                yld = self.fv(L.PO, r, 1)
                if r == L.PO_DIESEL_YIELDBACK_ROW:
                    continue
                if not out_code or not charge_code or not is_num(yld):
                    continue
                candidates = [ln for ln in by_unit_row[sec.unit]
                              if ln["code"] == charge_code]
                if r in L.PO_EXTRACT_DEEP_ROWS:
                    candidates = [ln for ln in by_unit_row[sec.unit]
                                  if ln["row"] == L.CS_EXTRACT_DEEP_ROW]
                if not candidates:
                    self.warn.add("yields",
                                  "Production-Out row {} ({} -> {}) has no charge line "
                                  "in unit {}".format(r, charge_code, out_code, sec.unit))
                    continue
                rules.append({"kind": "unit", "po_row": r, "unit": sec.unit,
                              "charge_line": candidates[0]["key"],
                              "charge_code": charge_code,
                              "out": out_code, "yield": float(yld)})

        # --- hydrotreater diesel yield-back
        #
        # The sheet multiplies each source row's *output* by (1/yield - 1), i.e.
        # charge x (1 - yield): the volume the solvent stream loses in
        # hydrotreating reappears as diesel charge. Applying the factor to the raw
        # charge instead inflates it by 1/yield.
        terms = []
        for r in L.PO_DIESEL_YIELDBACK_SOURCE_ROWS:
            yld = self.fv(L.PO, r, 1)
            if not is_num(yld) or not yld:
                continue
            terms.append({"source_po_row": r, "yield": float(yld)})
        out_code = norm_code(self.fv(L.PO, L.PO_DIESEL_YIELDBACK_ROW, 3))
        rules.append({"kind": "diesel_yield_back", "po_row": L.PO_DIESEL_YIELDBACK_ROW,
                      "out": out_code, "terms": terms})

        # --- transfers into pools
        for po_row, spec in L.PO_TRANSFER_ROWS.items():
            keys = [ln["key"] for ln in lines
                    if spec["rows"][0] <= ln["row"] <= spec["rows"][1]]
            rules.append({"kind": "transfer", "po_row": po_row,
                          "out": norm_code(spec["out"]),
                          "charge_lines": keys, "divisor": spec["divisor"]})
        return rules

    def _sheet_date_cols(self, sheet: str) -> List[Tuple[dt.date, int]]:
        if sheet not in self._date_cols_cache:
            self._date_cols_cache[sheet] = [
                (d, c) for c, d in read_date_row(self.wv[sheet], L.STREAM_DATE_ROW,
                                                 L.STREAM_FIRST_DATE_COL)]
        return self._date_cols_cache[sheet]

    # ------------------------------------------------------- formula resolution
    _VLOOKUP_PREFIX = r"VLOOKUP\(\s*\$?([A-Z]{1,3})\$?(\d+)\s*,\s*(?:\$?[A-Za-z ']*!?)?"
    _SUM_CELLS_RE = re.compile(r"SUM\(\s*((?:\$?[A-Z]{1,3}\$?\d+\s*,\s*){1,}\$?[A-Z]{1,3}\$?\d+)\s*\)")
    _CELL_RE = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)")

    def _parse_source(self, sheet: str, formula_text: str, cached: Any,
                      table: str) -> Dict[str, Any]:
        """Classify how a block sources a scalar (opening inventory / capacity).

        The workbook uses three different shapes for these and they are NOT
        interchangeable: summing two codes where the sheet looks up one would
        double the opening inventory of every block whose charge and sold codes
        are the same product.
        """
        from openpyxl.utils import column_index_from_string

        if formula_text:
            refs = re.findall(self._VLOOKUP_PREFIX + re.escape(table), formula_text)
            if refs:
                codes: List[str] = []
                for col, row in refs:
                    code = norm_code(self.fv(sheet, int(row),
                                             column_index_from_string(col)))
                    if code and code not in codes:
                        codes.append(code)
                out: Dict[str, Any] = {"kind": "codes", "codes": codes}
                # Some blocks share a product with another block and net the
                # other block's opening out of their own (e.g. MN's 9117 split).
                adj = [{"sign": -1 if s == "-" else 1, "row": int(r)}
                       for s, r in re.findall(r"([+-])\s*\$?G\$?(\d+)", formula_text)]
                if adj:
                    out["adjustments"] = adj
                return out

            m = self._SUM_CELLS_RE.search(formula_text)
            if m:
                rows = [int(r) for _, r in self._CELL_RE.findall(m.group(1))]
                return {"kind": "sum_of_block_openings", "rows": rows}

        if is_num(cached):
            return {"kind": "constant", "value": float(cached)}
        return {"kind": "codes", "codes": []}

    # ------------------------------------------------------- row classification
    #
    # The balance rows are NOT a uniform template. Blocks were hand-tweaked over
    # the years: some forecasts carry a split factor, the finished-diesel pool
    # aggregates its ten child blocks and applies a weekend rule, two LLN blocks
    # fold a downgrade into their production row, and Sonneborn sales come off the
    # charge schedule instead of open orders. Rather than hardcode each case we
    # classify every row's formula into a small set of terms, so the engine stays
    # generic and the per-product configuration becomes explicit data.

    _RX = {
        "prod_grid":   re.compile(r"'Production - Out'!\$C\$\d+"),
        "open_orders": re.compile(r"'Open Orders'!\$A:\$CH"),
        "blend":       re.compile(r"'Sales Forecast'!\$AH\$\d+"),
        "forecast":    re.compile(r"'Sales Forecast'!\$T:\$AF"),
        "bot_index":   re.compile(r"INDEX\(BotByMonth"),
        "bot_vlookup": re.compile(r"VLOOKUP\(\$E\d+,\s*BaseOilTransfer!\$A\$\d+:\$AR\$\d+,\s*(\d+)"),
        "charge_vl":   re.compile(r"VLOOKUP\(\$D\d+,\s*Charge,"),
        "to_diesel":   re.compile(r"VLOOKUP\(\$E\d+,\s*Charge_Gal_ToDiesel"),
        "hlookup_cs":  re.compile(r"HLOOKUP\([^,]+,\s*'Charge Schedule'!\$D\$(\d+):\$\w+\$\d+,\s*(\d+)"),
        "index_cs":    re.compile(r"INDEX\('Charge Schedule'!\$D\$(\d+):"),
        "vlookup_cs":  re.compile(r"VLOOKUP\(\$E(\d+),\s*'Charge Schedule'!\$B\$(\d+):\$\w+\$(\d+)"),
        "sum_cells":   re.compile(r"SUM\((\$?[A-Z]{1,3}\$?\d+(?:\s*,\s*\$?[A-Z]{1,3}\$?\d+)+)\)"),
        "weekday":     re.compile(r"IF\(WEEKDAY\([A-Z]{1,3}\d+,\s*(\d+)\)\s*>\s*(\d+)\s*,\s*(\d+)"),
        "factor":      re.compile(r"\*\s*\$D\$(\d+)\s*$"),
    }

    @staticmethod
    def _norm_formula(f: str) -> str:
        return re.sub(r"\$?[A-Z]{1,3}\$?\d+", "@", f) if f else ""

    def _row_regions(self, sheet: str, label: str, meta: Dict[str, Any],
                     date_cols: List[Tuple[dt.date, int]]) -> List[Dict[str, Any]]:
        """Split a balance row into contiguous spans that share a formula shape.

        The workbook's rows are not uniform across the horizon: some were filled
        from a later column so the early days are blank (planner-entered), some
        carry a special case only in the first column, and a few tail spans were
        pasted with a broken cross-sheet reference. Classifying per span keeps all
        of that explicit instead of picking one column and hoping.
        """
        regions: List[Dict[str, Any]] = []
        prev: Optional[str] = None
        for d, c in date_cols:
            f = self.ff(sheet, meta["row"], c)
            norm = self._norm_formula(f)
            if norm != prev:
                regions.append({"from": _iso(d), "formula": f})
                prev = norm
            regions[-1]["to"] = _iso(d)
        for reg in regions:
            reg["spec"] = self._classify_formula(sheet, label, meta, reg["formula"])
            reg["manual"] = not reg["formula"]
            del reg["formula"]
        if len(regions) > 1:
            self.warn.add("formula_regions",
                          "{} row {} ({}) changes formula {} time(s) across the "
                          "horizon at {}".format(
                              sheet, meta["row"], label, len(regions) - 1,
                              ", ".join(r["from"] for r in regions[1:])))
        return regions

    def _classify_formula(self, sheet: str, label: str, meta: Dict[str, Any],
                          f: str) -> Dict[str, Any]:
        """Turn one balance-row formula into a list of evaluatable terms."""
        in_code, out_code = meta["in"], meta["out"]
        terms: List[Dict[str, Any]] = []
        spec: Dict[str, Any] = {"terms": terms}

        if not f:
            spec["terms"] = [{"kind": "manual"}]
            return spec

        rx = self._RX

        m = rx["weekday"].search(f)
        if m:
            spec["weekday_override"] = {"basis": int(m.group(1)),
                                        "threshold": int(m.group(2)),
                                        "value": float(m.group(3))}

        m = rx["factor"].search(f)
        if m:
            v = self.fv(sheet, int(m.group(1)), L.STREAM_OUT_COL)
            spec["factor"] = float(v) if is_num(v) else 1.0

        m = rx["sum_cells"].search(f)
        if m:
            rows = [int(r) for _, r in self._CELL_RE.findall(m.group(1))]
            terms.append({"kind": "pool_sum", "sheet": sheet, "rows": rows})
            return spec

        if rx["prod_grid"].search(f):
            terms.append({"kind": "production_grid", "code": in_code})
        if rx["open_orders"].search(f):
            terms.append({"kind": "open_orders", "code": in_code})
        if rx["blend"].search(f):
            terms.append({"kind": "blend_demand", "code": in_code})
        if rx["forecast"].search(f):
            terms.append({"kind": "forecast_net", "code": in_code})
        if rx["bot_index"].search(f):
            terms.append({"kind": "base_oil_transfer_monthly", "code": in_code})
        m = rx["bot_vlookup"].search(f)
        if m:
            terms.append({"kind": "base_oil_transfer_column", "code": in_code,
                          "column": int(m.group(1))})
        if rx["charge_vl"].search(f):
            terms.append({"kind": "charge_first_match", "code": out_code})
        if rx["to_diesel"].search(f):
            terms.append({"kind": "to_diesel", "code": in_code})

        # Explicit Charge Schedule rows, reached three different ways.
        rows: List[int] = []
        for first, offset in rx["hlookup_cs"].findall(f):
            rows.append(int(first) + int(offset) - 1)
        for first in rx["index_cs"].findall(f):
            rows.append(int(first))
        for code_row, first, last in rx["vlookup_cs"].findall(f):
            want = norm_code(self.fv(sheet, int(code_row), L.STREAM_IN_COL))
            hit = next((r for r in range(int(first), int(last) + 1)
                        if norm_code(self.fv(L.CS, r, 2)) == want), None)
            if hit:
                rows.append(hit)
            else:
                self.warn.add("rows", "{} {}: code {} not found in Charge Schedule "
                              "rows {}-{}".format(sheet, label, want, first, last))
        for r in rows:
            terms.append({"kind": "charge_row", "row": r})

        if not terms:
            self.warn.add("rows", "{} row {} ({}): unrecognised formula {}"
                          .format(sheet, meta["row"], label, f[:120]))
            terms.append({"kind": "unknown"})
        return spec

    # ------------------------------------------------------------ block specs
    def block_specs(self) -> List[Dict[str, Any]]:
        specs: List[Dict[str, Any]] = []
        for sheet in L.STREAM_SHEETS:
            wsv = self.wv[sheet]
            r = 13
            while r <= wsv.max_row:
                label = text(wsv.cell(r, L.STREAM_LABEL_COL).value)
                if label != "Begin Inventory":
                    r += 1
                    continue
                block = self._parse_block(sheet, r)
                if block:
                    specs.append(block)
                    r = block["last_row"] + 1
                else:
                    r += 1

        # Resolve pooled openings (the finished-diesel block sums the opening cell
        # of each diesel sub-block) from sheet rows to block ids.
        by_open_row = {(b["sheet"], b["opening_row"]): b["id"] for b in specs}
        for b in specs:
            op = b["opening"]
            if op.get("kind") == "sum_of_block_openings":
                ids, missing = [], []
                for row in op["rows"]:
                    bid = by_open_row.get((b["sheet"], row))
                    (ids if bid else missing).append(bid or row)
                op["block_ids"] = ids
                if missing:
                    self.warn.add("blocks",
                                  "block {} pools rows {} that are not block openings"
                                  .format(b["id"], missing))
            for a in op.get("adjustments", []):
                a["block_id"] = by_open_row.get((b["sheet"], a["row"]))
                if not a["block_id"]:
                    self.warn.add("blocks",
                                  "block {} adjusts by row {}, which is not a block "
                                  "opening".format(b["id"], a["row"]))
        return specs

    def _parse_block(self, sheet: str, begin_row: int) -> Optional[Dict[str, Any]]:
        wsv = self.wv[sheet]
        rows: Dict[str, Dict[str, Any]] = {}
        order: List[str] = []
        r = begin_row
        last = begin_row
        date_cols = self._sheet_date_cols(sheet)
        while r <= min(begin_row + 20, wsv.max_row):
            label = text(wsv.cell(r, L.STREAM_LABEL_COL).value)
            if not label:
                break
            canon = label.strip()
            rows[canon] = {
                "row": r,
                "in": norm_code(wsv.cell(r, L.STREAM_IN_COL).value),
                "out": norm_code(wsv.cell(r, L.STREAM_OUT_COL).value),
                "formula": self.ff(sheet, r, L.STREAM_FIRST_DATE_COL),
                "formula_first": self.ff(sheet, r, L.STREAM_FIRST_DATE_COL),
                "regions": (self._row_regions(sheet, canon, {"row": r,
                            "in": norm_code(wsv.cell(r, L.STREAM_IN_COL).value),
                            "out": norm_code(wsv.cell(r, L.STREAM_OUT_COL).value)},
                            date_cols) if canon in _CLASSIFIED_LABELS else None),
            }
            order.append(canon)
            last = r
            if canon == "Excess Capacity":
                break
            r += 1
        if "End Inventory" not in rows and "End Inventory " not in rows:
            return None

        def get(name: str) -> Optional[Dict[str, Any]]:
            """Match a canonical label by prefix.

            Labels drift between blocks: LLN calls its production row
            "Production Out &DG DFO" because it folds the downgrade transfer into
            the same line, so an exact match would silently miss it.
            """
            for k in rows:
                if k.strip().startswith(name):
                    return rows[k]
            return None

        begin = get("Begin Inventory")
        end = get("End Inventory")
        cap = get("Tank Capacity")
        if not begin or not end:
            return None

        capacity = ({"kind": "codes", "codes": []} if not cap else
                    self._parse_source(sheet, cap.get("formula_first", ""),
                                       self.fv(sheet, cap["row"], L.STREAM_FIRST_DATE_COL),
                                       "Tank_Capacity"))
        open_row = begin_row - 1
        opening = self._parse_source(
            sheet, self.ff(sheet, open_row, L.STREAM_LABEL_COL),
            self.fv(sheet, open_row, L.STREAM_LABEL_COL), "tInv")

        prod_out = get("Production Out")
        extra_dg_offset = None
        if prod_out:
            m = re.search(r"HLOOKUP\([^,]+,\s*'Charge Schedule'!\$D\$(\d+):\$\w+\$(\d+),\s*(\d+)",
                          prod_out["formula"])
            if m:
                extra_dg_offset = int(m.group(1)) + int(m.group(3)) - 1

        blends = get("Blends")
        has_bot = bool(blends and "BotByMonth" in blends["formula"])

        charge_code = begin["in"]
        sold_code = begin["out"]
        # A row is planner-entered wherever it has no formula; some rows are manual
        # for part of the horizon only, so capture values for any row with at
        # least one manual span.
        manual = [k.strip() for k, v in rows.items()
                  if any(reg["manual"] for reg in (v.get("regions") or []))
                  or (v.get("regions") is None and not v["formula"])]
        return {
            "id": "{}:{}".format(sheet, begin_row),
            "sheet": sheet,
            "first_row": begin_row,
            "last_row": last,
            "charge_code": charge_code,
            "sold_code": sold_code,
            "label": text(wsv.cell(begin_row, 6).value) or charge_code,
            "rows": {k.strip(): {"row": v["row"], "in": v["in"], "out": v["out"],
                                 "regions": v.get("regions")}
                     for k, v in rows.items()},
            "row_order": [o.strip() for o in order],
            "capacity": capacity,
            "opening": opening,
            "opening_row": open_row,
            "prod_out_label": next((o.strip() for o in order
                                    if o.strip().startswith("Production Out")), None),
            "prod_out_extra_transfer_row": extra_dg_offset,
            "has_out_to_diesel": any(o.strip() == "Out to Diesel" for o in order),
            "blends_include_base_oil_transfer": has_bot,
            "manual_rows": manual,
        }

    def capacity_series(self, blocks: List[Dict[str, Any]], dates: List[dt.date]
                        ) -> Dict[str, Dict[str, float]]:
        """Tank capacity per block per day.

        Capacity is not constant: the sheet steps it up mid-horizon (e.g.
        `=CP23+300000`) when a tank project comes online. Capturing the series
        keeps that real planning data instead of freezing day-one capacity, and
        it is the shape the rebuild wants anyway (effective-dated capacity).
        Only blocks whose capacity actually varies are stored.
        """
        out: Dict[str, Dict[str, float]] = {}
        cols: Dict[str, Dict[dt.date, int]] = {}
        for sheet in L.STREAM_SHEETS:
            cols[sheet] = {d: c for c, d in
                           read_date_row(self.wv[sheet], L.STREAM_DATE_ROW,
                                         L.STREAM_FIRST_DATE_COL)}
        for b in blocks:
            cap = b["rows"].get("Tank Capacity")
            if not cap:
                continue
            series: Dict[str, float] = {}
            for d in dates:
                c = cols[b["sheet"]].get(d)
                if not c:
                    continue
                v = self.fv(b["sheet"], cap["row"], c)
                if is_num(v):
                    series[_iso(d)] = float(v)
            if len(set(series.values())) > 1:
                out[b["id"]] = series
                self.warn.add("capacity",
                              "block {} capacity changes over the horizon ({:,.0f} -> "
                              "{:,.0f}); modelled as an effective-dated series"
                              .format(b["id"], min(series.values()), max(series.values())))
        return out

    def manual_row_values(self, blocks: List[Dict[str, Any]],
                          dates: List[dt.date]) -> Dict[str, Dict[str, Dict[str, float]]]:
        """Cached values of planner-entered (formula-free) balance rows."""
        out: Dict[str, Dict[str, Dict[str, float]]] = {}
        cols: Dict[str, Dict[dt.date, int]] = {}
        for sheet in L.STREAM_SHEETS:
            cols[sheet] = {d: c for c, d in
                           read_date_row(self.wv[sheet], L.STREAM_DATE_ROW,
                                         L.STREAM_FIRST_DATE_COL)}
        for b in blocks:
            sheet = b["sheet"]
            per_label: Dict[str, Dict[str, float]] = {}
            for label in b["manual_rows"]:
                row = b["rows"][label]["row"]
                vals: Dict[str, float] = {}
                for d in dates:
                    c = cols[sheet].get(d)
                    if not c:
                        continue
                    v = self.fv(sheet, row, c)
                    if is_num(v) and v:
                        vals[_iso(d)] = float(v)
                if vals:
                    per_label[label] = vals
            if per_label:
                out[b["id"]] = per_label
        return out

    # ------------------------------------------------------------------- feeds
    def opening_inventory(self) -> Dict[str, float]:
        ws = self.wv[L.INVENTORY]
        out: Dict[str, float] = defaultdict(float)
        for r in range(L.INV_FIRST_ROW, ws.max_row + 1):
            code = norm_code(ws.cell(r, L.INV_CODE_COL).value)
            v = ws.cell(r, L.INV_GALLONS_COL).value
            if not code or not is_num(v):
                continue
            out[code] += float(v)
        return dict(out)

    def crude_opening_bbl(self) -> float:
        """Opening crude inventory in barrels, from the site summary block."""
        label = text(self.fv(L.INVENTORY, *L.INV_CRUDE_LABEL_CELL))
        v = self.fv(L.INVENTORY, *L.INV_CRUDE_TOTAL_CELL)
        if "total" not in label.lower():
            self.warn.add("crude_inventory",
                          "expected a crude total label at W13 but found {!r}; "
                          "crude opening inventory may be wrong".format(label))
        if not is_num(v):
            self.warn.add("crude_inventory", "no numeric crude total at X13")
            return 0.0
        return float(v)

    def open_orders(self) -> Dict[str, Dict[str, float]]:
        ws = self.wv[L.OPEN_ORDERS]
        dates = read_date_row(ws, L.OO_DATE_ROW, L.OO_FIRST_COL)
        if not dates:
            # row 6 mirrors row 8 with `=D8*1`; fall back to the text header row
            dates = read_date_row(ws, 8, L.OO_FIRST_COL)
        out: Dict[str, Dict[str, float]] = {}
        seen: Dict[str, int] = {}
        for r in range(L.OO_FIRST_ROW, min(L.OO_LAST_ROW, ws.max_row) + 1):
            code = norm_code(ws.cell(r, L.OO_CODE_COL).value)
            if not code:
                continue
            if code in seen:
                self.warn.add("open_orders",
                              "duplicate rows for {} (rows {} and {}); the workbook's "
                              "VLOOKUP only sees the first".format(code, seen[code], r))
                continue
            seen[code] = r
            per_date: Dict[str, float] = {}
            for c, d in dates:
                v = ws.cell(r, c).value
                if is_num(v) and v:
                    # feed is negative gallons; store as positive demand
                    per_date[_iso(d)] = -float(v)
            out[code] = per_date
        return out

    def _monthly_grid(self, sheet: str, hdr_row: int, code_col: int,
                      first_col: int, first_row: int, last_row: int
                      ) -> Dict[str, Dict[str, float]]:
        months = self._month_header(sheet, hdr_row, first_col)
        out: Dict[str, Dict[str, float]] = {}
        for r in range(first_row, last_row + 1):
            code = norm_code(self.fv(sheet, r, code_col))
            if not code:
                continue
            vals = {}
            for i, m in enumerate(months):
                v = self.fv(sheet, r, first_col + i)
                if is_num(v) and v:
                    vals[m] = float(v)
            if vals:
                out.setdefault(code, vals)
        return out

    def sales_forecast(self) -> Dict[str, Dict[str, float]]:
        return self._monthly_grid(L.SALES_FORECAST, L.SF_HEADER_ROW,
                                  L.SF_RATE_CODE_COL, L.SF_RATE_FIRST_COL,
                                  L.SF_FIRST_ROW, L.SF_LAST_ROW)

    def blend_component_demand(self) -> Dict[str, Dict[str, float]]:
        return self._monthly_grid(L.SALES_FORECAST, L.SF_HEADER_ROW,
                                  L.SF_BLEND_CODE_COL, L.SF_BLEND_FIRST_COL,
                                  L.SF_FIRST_ROW, L.SF_LAST_ROW)

    def base_oil_transfer(self) -> Dict[str, Dict[str, float]]:
        """Monthly base-oil transfer volumes, keyed by the month header in AW:BH
        but read from the BotByMonth value range that the sheet actually INDEXes."""
        months = self._month_header(L.BASE_OIL_TRANSFER, L.BOT_MONTH_HDR_ROW,
                                    L.BOT_HEADER_FIRST_COL)
        out: Dict[str, Dict[str, float]] = {}
        for r in range(L.BOT_FIRST_ROW, L.BOT_LAST_ROW + 1):
            code = norm_code(self.fv(L.BASE_OIL_TRANSFER, r, L.BOT_CODE_COL))
            if not code:
                continue
            vals = {}
            for i, m in enumerate(months):
                v = self.fv(L.BASE_OIL_TRANSFER, r, L.BOT_FIRST_MONTH_COL + i)
                if is_num(v) and v:
                    vals[m] = float(v)
            if vals:
                out.setdefault(code, vals)
        return out

    def base_oil_transfer_columns(self) -> Dict[str, Dict[str, float]]:
        """The A:AR analysis table, addressed by absolute column index."""
        out: Dict[str, Dict[str, float]] = {}
        for r in range(L.BOT_TABLE_FIRST_ROW, L.BOT_TABLE_LAST_ROW + 1):
            code = norm_code(self.fv(L.BASE_OIL_TRANSFER, r, L.BOT_TABLE_CODE_COL))
            if not code:
                continue
            vals = {}
            for c in range(2, 45):
                v = self.fv(L.BASE_OIL_TRANSFER, r, c)
                if is_num(v):
                    vals[str(c)] = float(v)
            out.setdefault(code, vals)
        return out

    def daily_grid_dates(self) -> Tuple[dt.date, List[dt.date]]:
        ws = self.wv[L.STREAM_SHEETS[0]]
        as_of = as_date(ws.cell(*L.STREAM_AS_OF_CELL).value)
        dates = [d for _, d in read_date_row(ws, L.STREAM_DATE_ROW,
                                             L.STREAM_FIRST_DATE_COL)]
        return as_of, dates

    # -------------------------------------------------------------------- run
    def run(self, outdir: str) -> Dict[str, str]:
        self.validate_layout()
        as_of, dates = self.daily_grid_dates()
        lines = self.charge_lines()
        cdates = self.charge_dates()
        cy, loss = self.crude_yields()
        blocks = self.block_specs()

        reference = {
            "meta": {
                "source_file": os.path.basename(self.path),
                "imported_at": dt.datetime.now().isoformat(timespec="seconds"),
                "gal_per_bbl": GAL_PER_BBL,
            },
            "products": self.products(),
            "tank_capacity": self.tank_capacity(),
            "inventory_targets": self.inventory_targets(),
            "crude_yields": cy,
            "crude_loss_fraction": loss,
            "run_rates": self.run_rates(),
            "other_rates": self.other_rates(),
            "crude_plan": self.crude_plan(),
            "blend_bom": self.blend_bom(),
            "charge_lines": lines,
            "yield_rules": self.yield_rules(lines),
            "blocks": blocks,
        }

        scenario = {
            "meta": {
                "name": "workbook-as-imported",
                "as_of": _iso(as_of) if as_of else None,
                "source_file": os.path.basename(self.path),
            },
            "dates": [_iso(d) for d in dates],
            "opening_inventory": self.opening_inventory(),
            "crude_opening_bbl": self.crude_opening_bbl(),
            "open_orders": self.open_orders(),
            "sales_forecast": self.sales_forecast(),
            "blend_component_demand": self.blend_component_demand(),
            "base_oil_transfer": self.base_oil_transfer(),
            "base_oil_transfer_columns": self.base_oil_transfer_columns(),
            "charge_schedule": self.charge_schedule(lines, cdates),
            "manual_block_rows": self.manual_row_values(blocks, dates),
            "capacity_series": self.capacity_series(blocks, dates),
        }

        os.makedirs(outdir, exist_ok=True)
        paths = {}
        for name, payload in (("reference.json", reference),
                              ("scenario.json", scenario)):
            p = os.path.join(outdir, name)
            with open(p, "w", encoding="utf-8") as f:
                json.dump(payload, f, indent=1, sort_keys=False)
            paths[name] = p
        if self.warn.items:
            p = os.path.join(outdir, "import-warnings.txt")
            with open(p, "w", encoding="utf-8") as f:
                f.write("\n".join(self.warn.items))
            paths["warnings"] = p
        return paths


def import_workbook(path: str, outdir: str) -> Dict[str, str]:
    return WorkbookImporter(path).run(outdir)
