"""API and persistence tests.

Covers the three things Phase 1 has to get right: overrides survive a re-sync,
freezing a scenario reproduces the engine's numbers, and every endpoint the
front end calls actually exists.
"""
import datetime as dt
import os
import re
import sys

import pytest

ROOT = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
sys.path.insert(0, os.path.join(ROOT, "src"))

SEED = os.path.join(ROOT, "data", "seed")
pytestmark = pytest.mark.skipif(
    not os.path.exists(os.path.join(SEED, "reference.json")),
    reason="seed data not built; run scripts/seed.py")


@pytest.fixture(scope="module")
def db(tmp_path_factory):
    """A throwaway database seeded exactly the way scripts/init_db.py does."""
    path = tmp_path_factory.mktemp("db") / "test.db"
    os.environ["DATABASE_URL"] = "sqlite:///" + str(path)
    # An empty upload folder, so these tests read the workbook seed whatever a
    # planner happens to have uploaded on this machine. Several of them assert
    # that a scenario frozen from staging matches `scenario.json` exactly, which
    # only means anything when staging came from that file.
    os.environ["INVPLANNER_UPLOAD_DIR"] = str(tmp_path_factory.mktemp("uploads"))
    for mod in [m for m in list(sys.modules) if m.startswith("invplanner.db")]:
        del sys.modules[mod]

    from invplanner.db.session import SessionLocal, init_db
    from invplanner.db import service as svc

    init_db()
    session = SessionLocal()
    svc.load_reference_document(session, os.path.join(SEED, "reference.json"))
    session.commit()
    svc.sync_all(session, actor="test")
    yield session
    session.close()


def test_sync_populates_staging_from_source(db):
    from invplanner.db.models import Feed, StagingCell
    n = db.query(StagingCell).filter(StagingCell.feed == Feed.INVENTORY).count()
    assert n > 100


def test_override_survives_resync_and_is_flagged_when_source_moves(db):
    """The whole point of the staging layer: a re-sync must not clobber a
    planner's edit, and it must show when the source has moved underneath it."""
    from invplanner.db import service as svc
    from invplanner.db.connectors import default_connectors
    from invplanner.db.models import Feed, StagingCell

    cell = (db.query(StagingCell)
            .filter(StagingCell.feed == Feed.INVENTORY)
            .order_by(StagingCell.k1).first())
    original_source = cell.source_value

    svc.set_override(db, cell.id, 999999.0, actor="tester", reason="unit test")
    assert cell.effective_value == 999999.0
    assert not cell.is_stale

    conn = next(c for c in default_connectors(svc.SEED_DIR) if c.feed == Feed.INVENTORY)
    svc.sync_feed(db, conn, actor="test")
    db.refresh(cell)

    assert cell.override_value == 999999.0, "re-sync clobbered the override"
    assert cell.source_value == original_source
    assert cell.effective_value == 999999.0

    # Now move the source and confirm the override is reported as stale.
    cell.source_value = (original_source or 0.0) + 5000.0
    db.commit()
    db.refresh(cell)
    assert cell.is_stale

    svc.set_override(db, cell.id, None, actor="tester")
    db.refresh(cell)
    assert cell.override_value is None
    assert not cell.is_stale


def test_scenario_freeze_reproduces_engine_numbers(db):
    """A scenario built from staging must simulate identically to the JSON
    scenario the parity harness validated against the workbook."""
    import json

    from invplanner.db import service as svc
    from invplanner.engine import Reference, Scenario, simulate

    with open(os.path.join(SEED, "scenario.json"), encoding="utf-8") as f:
        seed_scn = json.load(f)
    as_of = seed_scn["meta"]["as_of"]

    scenario = svc.create_scenario(db, "test", svc._d(as_of), horizon_days=90,
                                   actor="test")
    db_sim = svc.simulate_scenario(db, scenario)

    ref = Reference.load(os.path.join(SEED, "reference.json"))
    json_sim = simulate(ref, Scenario.load(os.path.join(SEED, "scenario.json")))

    dates = [svc._d(x) for x in scenario.inputs["dates"]][:60]
    checked = 0
    for block in ref.blocks:
        a = db_sim.balances[block["id"]]["End Inventory"]
        b = json_sim.balances[block["id"]]["End Inventory"]
        for d in dates:
            assert a[d] == pytest.approx(b[d], rel=1e-9, abs=1e-6), \
                "{} diverges on {}".format(block["id"], d)
            checked += 1
    assert checked > 1000


def test_schedule_edit_changes_the_projection(db):
    from invplanner.db import service as svc
    from invplanner.db.models import Scenario as ScenarioModel

    scenario = db.query(ScenarioModel).order_by(ScenarioModel.id.desc()).first()
    ref = svc.engine_reference(db, scenario.reference_version_id)
    block = next(b for b in ref.blocks if b["id"] == "MN:14")
    day = svc._d(scenario.inputs["dates"][3])

    before = svc.simulate_scenario(db, scenario).balances[block["id"]]["End Inventory"][day]

    line = next(l for l in ref.charge_lines if l["unit"] == "MEK")
    from invplanner.db.models import ScheduleEntry
    entry = (db.query(ScheduleEntry)
             .filter(ScheduleEntry.scenario_id == scenario.id,
                     ScheduleEntry.line_key == line["key"],
                     ScheduleEntry.date == day).first())
    if entry is None:
        entry = ScheduleEntry(scenario_id=scenario.id, line_key=line["key"],
                              date=day, bbl=0.0)
        db.add(entry)
    entry.bbl = (entry.bbl or 0.0) + 2500.0
    scenario.schedule_version += 1
    db.commit()
    svc.invalidate(scenario.id)

    after = svc.simulate_scenario(db, scenario).balances[block["id"]]["End Inventory"][day]
    assert after != before, "charging more feed left the projection unchanged"


def test_stream_sheet_returns_every_product_with_full_balance(db):
    """The stream view must reproduce the workbook tab: all products on the
    sheet, each with its balance rows, over a date window."""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db.models import Scenario as ScenarioModel
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    sid = db.query(ScenarioModel).order_by(ScenarioModel.id).first().id

    streams = client.get("/api/scenarios/{}/streams".format(sid)).json()
    sheets = {s["sheet"] for s in streams}
    assert {"Solvents", "LLN", "Cstock", "MN", "HN", "Napthas", "CRUDE"} <= sheets

    r = client.get("/api/scenarios/{}/stream".format(sid),
                   params={"sheet": "Solvents", "days": 5})
    assert r.status_code == 200
    d = r.json()
    assert len(d["dates"]) == 5
    assert len(d["blocks"]) == 18

    block = d["blocks"][0]
    measures = {row["measure"] for row in block["rows"]}
    assert {"Begin Inventory", "Production In", "Sales", "Forecast",
            "End Inventory", "Excess Capacity"} <= measures
    for row in block["rows"]:
        assert len(row["values"]) == 5

    # The balance must actually balance on every day shown.
    by = {row["measure"]: row["values"] for row in block["rows"]}
    for i in range(5):
        expected_end = (by["Begin Inventory"][i] + by["Production In"][i]
                        - by["Sales"][i] - by["Forecast"][i] - by["Blends"][i]
                        - by["Production Out"][i] - by.get("Out to Diesel", [0] * 5)[i])
        assert by["End Inventory"][i] == pytest.approx(expected_end, abs=1e-6)

    # Crude is a different shape: barrels, and receipts instead of production in.
    c = client.get("/api/scenarios/{}/stream".format(sid),
                   params={"sheet": "CRUDE", "days": 3}).json()
    assert c["unit"] == "bbl"
    crude = {row["measure"]: row["values"] for row in c["blocks"][0]["rows"]}
    assert "Receipts" in crude
    assert crude["Begin Inventory"][0] > 0, "crude must open from the site summary"

    app.dependency_overrides.clear()


def test_dashboard_groups_cover_every_live_product_exactly_once(db):
    """No product may vanish from the dashboards, and none may be double-counted.

    The "other" group is the safety net: a newly added block lands there rather
    than disappearing, so this stays true as the product list changes.

    Retired products are the one exception, and they are checked separately
    below rather than merely excused here - the whole risk of hiding something
    is that it comes back through a route nobody was looking at.
    """
    from fastapi.testclient import TestClient

    from invplanner import groups
    from invplanner import model_config as cfg
    from invplanner.api.main import app
    from invplanner.db import service as svc
    from invplanner.db.models import Scenario as ScenarioModel
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    scenario = db.query(ScenarioModel).order_by(ScenarioModel.id).first()
    sid = scenario.id

    ref = svc.engine_reference(db, scenario.reference_version_id)
    expected = {b["id"] for b in cfg.live_blocks(ref.blocks)} | {"CRUDE:14"}
    retired = {b["id"] for b in ref.blocks
               if cfg.is_retired(code=b.get("charge_code"))}
    assert retired, "the fixture must still contain something retired to hide"

    seen = []
    for g in client.get("/api/dashboards").json():
        d = client.get("/api/scenarios/{}/dashboard".format(sid),
                       params={"group": g["id"], "days": 60}).json()
        for section in d["sections"]:
            for p in section["products"]:
                seen.append(p["block_id"])

    assert len(seen) == len(set(seen)), "a product appears in more than one group"
    assert set(seen) == expected, "groups do not cover every product: {}".format(
        expected.symmetric_difference(seen))
    assert not (retired & set(seen)), (
        "retired products are back on the dashboard: {}"
        .format(sorted(retired & set(seen))))

    # Series must be downsampled but keep the true extremes and the last day.
    d = client.get("/api/scenarios/{}/dashboard".format(sid),
                   params={"group": "solvents", "days": 366, "max_points": 90}).json()
    p = d["sections"][0]["products"][0]
    assert len(p["end"]) <= 95
    assert p["dates"][-1] == d["to"]
    assert p["peak"] >= max(p["end"]) and p["trough"] <= min(p["end"])

    app.dependency_overrides.clear()


def test_planned_downtime_round_trips_and_splits(db):
    """Planners set downtime before the optimizer runs. Clearing a day inside a
    window must split it, not delete the whole turnaround."""
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import service as svc
    from invplanner.db.models import Scenario as ScenarioModel
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    scenario = db.query(ScenarioModel).order_by(ScenarioModel.id).first()
    sid = scenario.id
    before_version = scenario.schedule_version

    r = client.post("/api/scenarios/{}/downtime".format(sid), json={
        "unit": "MEK", "start": "2026-09-10", "end": "2026-09-14",
        "reason": "test turnaround"})
    assert r.status_code == 200 and r.json()["days"] == 5

    days = svc.downtime_days(db, sid)["MEK"]
    for day in range(10, 15):
        assert dt.date(2026, 9, day) in days

    # editing downtime must invalidate the cached simulation
    db.refresh(scenario)
    assert scenario.schedule_version > before_version

    # clearing a day in the middle splits the window in two
    t = client.post("/api/scenarios/{}/downtime/toggle".format(sid),
                    json={"unit": "MEK", "date": "2026-09-12"}).json()
    assert t["down"] is False
    wins = [x for x in client.get("/api/scenarios/{}/downtime".format(sid)).json()
            ["windows"] if x["unit"] == "MEK" and x["start"] >= "2026-09-10"]
    assert sorted((x["start"], x["end"]) for x in wins) == [
        ("2026-09-10", "2026-09-11"), ("2026-09-13", "2026-09-14")]

    # the schedule grid reports the days so the UI can lock those cells
    sched = client.get("/api/scenarios/{}/schedule".format(sid),
                       params={"days": 60}).json()
    mek = next(u for u in sched["units"] if u["unit"] == "MEK")
    assert "2026-09-11" in mek["downtime"]
    assert "2026-09-12" not in mek["downtime"]

    for wnd in wins:
        client.delete("/api/scenarios/{}/downtime/{}".format(sid, wnd["id"]))

    app.dependency_overrides.clear()


def test_reference_data_is_editable_and_changes_the_projection(db):
    """Yields, rates and capacities drive every downstream number, so they must
    be adjustable - and an edit has to reach the projection, not just the table."""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import service as svc
    from invplanner.db.models import Scenario as ScenarioModel
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    scenario = db.query(ScenarioModel).order_by(ScenarioModel.id).first()
    if scenario is None:      # not dependent on another test having run first
        import datetime as dt
        scenario = svc.create_scenario(db, "reference-edit", dt.date(2026, 7, 23),
                                       horizon_days=42, actor="test")

    groups = {g["kind"]: g for g in client.get("/api/reference").json()["groups"]}
    assert groups["unit_yield"]["count"] > 0
    assert groups["crude_yield"]["count"] > 0

    rows = client.get("/api/reference/unit_yield").json()["rows"]

    # pick a yield that actually produces something in the window, so the edit
    # has an observable effect
    sim = svc.simulate_scenario(db, scenario)
    dates = [svc._d(x) for x in scenario.inputs["dates"]][:60]
    target = day = product = None
    for row in rows:
        _, _, out = row["k1"].split("|")
        hit = next((d for d in dates if sim.prod(out, d) > 0), None)
        if hit is not None:
            target, day, product = row, hit, out
            break
    assert target is not None, "no unit yield produces anything in the window"

    before = sim.prod(product, day)
    assert before > 0

    r = client.patch("/api/reference/unit_yield", json={
        "k1": target["k1"], "value": round(target["source_value"] * 0.5, 8),
        "reason": "test"})
    assert r.status_code == 200 and r.json()["edited"] is True

    db.refresh(scenario)
    after = svc.simulate_scenario(db, scenario).prod(product, day)
    assert after < before, "halving a yield must reduce what the unit makes"

    # the imported value is kept alongside the edit
    row = next(x for x in client.get("/api/reference/unit_yield").json()["rows"]
               if x["k1"] == target["k1"])
    assert row["source_value"] == pytest.approx(target["source_value"])
    assert row["edited"] is True

    # clearing it restores the imported value
    client.patch("/api/reference/unit_yield",
                 json={"k1": target["k1"], "value": None})
    db.refresh(scenario)
    restored = svc.simulate_scenario(db, scenario).prod(product, day)
    assert restored == pytest.approx(before, rel=1e-9)

    app.dependency_overrides.clear()


def test_a_retired_product_is_gone_from_every_screen(db):
    """Retired means retired: not on a chart, not in a picker, not editable.

    The solver has dropped 9202 and 4325 since `RETIRED` was written, but the
    planner side kept showing them - a chart with no production, a tank
    capacity nobody could reach, and charge rows the next run would silently
    discard. Hiding them is spread across six endpoints and the reference
    editor, so this asserts the outcome rather than the mechanism: name a
    retired product and it must not appear anywhere a planner looks.

    Deliberately not asserted: that the blocks are gone from the *simulation*.
    Nothing needs them - deleting both from the reference moves no row anywhere
    - but the engine's job is to reproduce the workbook cell for cell so the
    parity harness has something to compare against, and a block it does not
    carry is a block parity stops checking. Retiring happens a layer up.
    """
    import json as _json

    from fastapi.testclient import TestClient

    from invplanner import model_config as cfg
    from invplanner.api.main import app
    from invplanner.db import reference_edit as refedit
    from invplanner.db import service as svc
    from invplanner.db.models import RefKind, ReferenceVersion
    from invplanner.db.models import Scenario as ScenarioModel
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    scenario = db.query(ScenarioModel).order_by(ScenarioModel.id).first()
    sid = scenario.id
    ref = svc.engine_reference(db, scenario.reference_version_id)

    dead_codes = sorted(cfg.RETIRED["products"])
    assert dead_codes, "nothing is retired, so this test proves nothing"
    dead_blocks = {b["id"] for b in ref.blocks
                   if b.get("charge_code") in dead_codes}
    dead_lines = {l["key"] for l in ref.charge_lines
                  if l["key"] not in {x["key"]
                                      for x in cfg.live_charge_lines(ref.charge_lines)}}
    assert dead_blocks and dead_lines

    # 1. the product picker and the stream sheets
    codes = {p["code"] for p in
             client.get("/api/scenarios/{}/products".format(sid)).json()}
    assert not (set(dead_codes) & codes)

    for sheet in {b["sheet"] for b in ref.blocks if b["id"] in dead_blocks}:
        d = client.get("/api/scenarios/{}/stream".format(sid),
                       params={"sheet": sheet, "days": 3}).json()
        assert not (dead_blocks & {b["block_id"] for b in d["blocks"]})

    # 2. the charts, in every family including "other"
    for g in client.get("/api/dashboards").json():
        d = client.get("/api/scenarios/{}/dashboard".format(sid),
                       params={"group": g["id"], "days": 60}).json()
        shown = {p["block_id"] for s in d["sections"] for p in s["products"]}
        assert not (dead_blocks & shown), g["id"]

    # 3. the alerts, which is the one screen that reports by exception
    a = client.get("/api/scenarios/{}/alerts".format(sid),
                   params={"days": 120}).json()
    assert not (dead_blocks & {x["block_id"] for x in a["alerts"]})

    # 4. the charge grid, and the export taken from it
    sch = client.get("/api/scenarios/{}/schedule".format(sid),
                     params={"days": 7}).json()
    on_grid = {l["key"] for u in sch["units"] for l in u["lines"]}
    assert not (dead_lines & on_grid)
    assert "TOLLING" not in {u["unit"] for u in sch["units"]}

    # 5. a direct hit on a hidden thing is refused, not quietly served
    for block_id in dead_blocks:
        r = client.get("/api/scenarios/{}/projection".format(sid),
                       params={"block_id": block_id})
        assert r.status_code == 404, block_id
    for key in dead_lines:
        r = client.patch("/api/scenarios/{}/schedule".format(sid),
                         json={"line_key": key, "date": "2026-07-23", "bbl": 1.0})
        assert r.status_code == 400, key

    # 6. the reference editor, where a value that changes nothing is worse than
    #    no value at all
    rv = db.query(ReferenceVersion).get(scenario.reference_version_id)
    for kind in RefKind.ALL:
        rows = refedit.list_values(db, rv, kind)
        items = rows["rows"] if isinstance(rows, dict) else rows
        for r in items:
            key = "{}|{}".format(r.get("k1"), r.get("k2"))
            assert not any(c in key.split("|") or c in str(r.get("k1")).split("|")
                           for c in dead_codes), (kind, r.get("k1"), r.get("k2"))

    # 7. and the solver, which is where this all started
    from invplanner.engine import Reference, Scenario, simulate
    from invplanner.modelprep import build
    eref = Reference.load(os.path.join(SEED, "reference.json"))
    escn = Scenario.load(os.path.join(SEED, "scenario.json"))
    spec = build(eref, escn, simulate(eref, escn, physical=True),
                 horizon=escn.dates[:21])
    assert not (set(dead_codes) & set(spec.products))
    assert not (set(dead_codes) & {l.get("code") for l in spec.charge_lines})
    dropped = _json.dumps(spec.dropped)
    for c in dead_codes:
        assert c in dropped, "{} must be dropped on the record, not silently".format(c)

    app.dependency_overrides.clear()


def test_every_endpoint_the_ui_calls_exists():
    """Guards against the front end drifting away from the API."""
    from invplanner.api.main import app

    js_path = os.path.join(ROOT, "src", "invplanner", "web", "app.js")
    with open(js_path, encoding="utf-8") as f:
        js = f.read()

    called = set(re.findall(r"""api\(\s*[`'"]([^`'"]+)""", js))
    assert called, "no api() calls found - did the front end change shape?"

    patterns = []
    for route in app.routes:
        if not hasattr(route, "path"):
            continue
        patterns.append(re.compile(
            "^" + re.sub(r"\{[^}]+\}", "[^/]+", route.path) + "$"))

    unresolved = []
    for call in sorted(called):
        path = call.split("?")[0]
        path = re.sub(r"\$\{[^}]*\}", "X", path)
        if not any(p.match(path) for p in patterns):
            unresolved.append(call)
    assert not unresolved, "front end calls endpoints that do not exist: {}".format(
        unresolved)


def test_schedule_exports_as_paste_blocks(db):
    """The export hands over blocks to paste, not a rewritten workbook.

    One block per unit with its own target cell, because the rows between the
    unit blocks carry a total, the date headers and the "enter BBLs in yellow
    cells" note - a single contiguous paste would flatten all of them.
    """
    import io

    from fastapi.testclient import TestClient
    from openpyxl import load_workbook

    from invplanner.api.main import app
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    scenarios = client.get("/api/scenarios").json()
    sid = scenarios[0]["id"]
    r = client.get("/api/scenarios/{}/schedule.xlsx?days=14".format(sid))
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    assert "attachment" in r.headers["content-disposition"]

    ws = load_workbook(io.BytesIO(r.content))["Paste blocks"]
    headers = [ws.cell(row, 1).value for row in range(1, ws.max_row + 1)
               if "paste at" in str(ws.cell(row, 1).value or "")]
    assert headers, "no paste blocks in the export"
    # every block names a cell, and no two units share one
    targets = [h.split("paste at ")[1] for h in headers]
    assert len(targets) == len(set(targets))
    app.dependency_overrides.clear()


def test_export_targets_the_column_the_date_really_sits_in(db):
    """The paste target is computed from a layout anchor, so the anchor has to be
    right. Checked against the workbook when it is present - a wrong column is
    the failure that looks exactly like success."""
    import datetime as dt
    import os

    from invplanner import exporter, layout

    y, m, d = (int(x) for x in layout.CS_GRID_FIRST_DATE.split("-"))
    assert exporter.grid_column(dt.date(y, m, d)) == layout.CS_FIRST_COL

    book = os.path.join(ROOT, "..", layout.WORKBOOK_DEFAULT)
    if not os.path.exists(book):
        pytest.skip("workbook not present")
    from openpyxl import load_workbook
    ws = load_workbook(book, data_only=True, read_only=True)["Charge Schedule"]
    row = next(ws.iter_rows(min_row=layout.CS_DATE_ROW_BBL,
                            max_row=layout.CS_DATE_ROW_BBL,
                            min_col=1, max_col=200, values_only=True))
    for col, cell in enumerate(row, start=1):
        if isinstance(cell, dt.datetime):
            assert exporter.grid_column(cell.date()) == col, (
                "{} sits in column {}, the export would target {}".format(
                    cell.date(), col, exporter.grid_column(cell.date())))
            break


def test_a_result_scenario_never_charges_a_unit_that_is_down(db):
    """The optimizer creates no charge variable on an outage day, so it returns
    nothing for one - and `_write_result_scenario` copies the base grid before
    it writes. Without the clear, the planner's own charge survives into the
    result and the schedule shows a unit running while it is down. It reached
    the verifier as an unexplained feed shortfall, which is the symptom rather
    than the cause.
    """
    import datetime as dt

    from invplanner.db import optimizer_service as optsvc
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    base = svc.create_scenario(db, "downtime write-back", dt.date(2026, 7, 23),
                               60, "test")
    line = (db.query(ScheduleEntry)
            .filter(ScheduleEntry.scenario_id == base.id,
                    ScheduleEntry.bbl != 0)
            .order_by(ScheduleEntry.date).first())
    assert line is not None, "seeded scenario has no charge to inherit"
    unit = line.line_key.split("#")[0]
    svc.add_downtime(db, base.id, unit, line.date, line.date, "test outage")

    class _Result:                       # what a solver hands back
        charge = {}                      # nothing on the outage day, as in life
        transfer_bbl = {}

    new_id = optsvc._write_result_scenario(db, base, _Result(), "test", 0)

    kept = (db.query(ScheduleEntry)
            .filter(ScheduleEntry.scenario_id == new_id,
                    ScheduleEntry.line_key == line.line_key,
                    ScheduleEntry.date == line.date).one())
    assert kept.bbl == 0.0, (
        "{} charge on {} survived into the result".format(
            line.line_key, line.date))

    # And the verifier names it rather than leaving it to surface downstream as
    # an unexplained feed shortfall. Checked against the *base* scenario, which
    # still carries the charge the result no longer does - so this asserts the
    # check works, not merely that the fix removed its input.
    from invplanner.optimizer import verify

    down = svc.downtime_days(db, base.id)
    assert down.get(unit), "downtime did not register"
    ref = svc.engine_reference(db, base.reference_version_id)
    dates = [svc._d(x) for x in base.inputs["dates"]][:30]
    claimed = {"lost_sales_gal": 0.0, "downgrade_gal": 0.0}

    blind = verify.verify(ref, svc.engine_scenario(db, base), claimed, dates)
    assert blind["respects_downtime"], "no downtime passed, nothing to report"

    seen = verify.verify(ref, svc.engine_scenario(db, base), claimed, dates,
                         downtime=down)
    assert not seen["respects_downtime"]
    assert not seen["verified"]
    assert any(r["line"] == line.line_key for r in seen["charged_when_down"])
    assert "the unit is down" in seen["note"]


# ----------------------------------------------------------- schedule fill
def _fill_scenario(db, name):
    """A scenario of its own, so a bulk edit cannot disturb its neighbours."""
    from invplanner.db import service as svc
    from invplanner import model_config as cfg

    import json
    with open(os.path.join(SEED, "scenario.json"), encoding="utf-8") as f:
        as_of = json.load(f)["meta"]["as_of"]
    scenario = svc.create_scenario(db, name, svc._d(as_of), horizon_days=40,
                                   actor="test")
    ref = svc.engine_reference(db, scenario.reference_version_id)
    live = {l["key"] for l in cfg.live_charge_lines(ref.charge_lines)}
    dates = [svc._d(x) for x in scenario.inputs["dates"]]
    return scenario, live, dates


def _set_cell(db, scenario_id, line_key, day, bbl):
    """Set one cell, whether or not a row is already there."""
    from invplanner.db.models import ScheduleEntry
    row = (db.query(ScheduleEntry)
           .filter(ScheduleEntry.scenario_id == scenario_id,
                   ScheduleEntry.line_key == line_key,
                   ScheduleEntry.date == day).first())
    if row is None:
        row = ScheduleEntry(scenario_id=scenario_id, line_key=line_key,
                            date=day, bbl=0.0)
        db.add(row)
    row.bbl = bbl
    db.commit()


def test_filling_a_rate_writes_the_window_and_clears_the_siblings(db):
    """The cold-start setup: one rate across every day, one line at a time.

    The sibling clear is the part that is easy to leave out and expensive to
    leave out. A unit runs one feed at a time, so a fill that writes only its own
    line leaves whatever was already scheduled on the others and doubles the
    unit's charge on those days - a schedule the plant cannot run, handed to the
    optimizer as though a planner meant it.
    """
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    scenario, live, dates = _fill_scenario(db, "fill-rate")
    rose = sorted(k for k in live if k.startswith("ROSE#"))
    assert len(rose) > 1, "this test needs a unit with more than one line"
    target, sibling = rose[0], rose[1]
    down = svc.downtime_days(db, scenario.id).get("ROSE", set())
    window = [d for d in dates[:20] if d not in down]
    assert len(window) > 5

    svc.fill_schedule(db, scenario.id, "ROSE", 0.0, live_keys=live, actor="test")
    _set_cell(db, scenario.id, sibling, window[5], 1234.0)

    r = svc.fill_schedule(db, scenario.id, target, 3000.0,
                          start=window[0], end=window[-1], live_keys=live,
                          actor="test")

    got = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == target)}
    for d in window:
        assert got.get(d) == 3000.0, "no rate written on {}".format(d)
    assert not got.get(dates[-1]), "wrote past the end date"

    sib = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == sibling)}
    assert sib.get(window[5]) == 0.0
    assert r["siblings_cleared"] == 1


def test_filling_a_whole_unit_with_a_rate_is_refused(db):
    """Most units carry several lines. Writing one rate onto all of them runs
    every line at once, which is not what anyone means by "run MEK at 3,000"."""
    import pytest
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    scenario, live, dates = _fill_scenario(db, "fill-unit")
    window = set(dates)
    with pytest.raises(ValueError) as e:
        svc.fill_schedule(db, scenario.id, "MEK", 3000.0, live_keys=live,
                          actor="test")
    assert "single line" in str(e.value)

    # Clearing it is the operation that does make sense, and handing the
    # optimizer an empty schedule to decide for is the whole of a cold start.
    svc.fill_schedule(db, scenario.id, "MEK", 0.0, live_keys=live, actor="test")
    left = [e.line_key for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id)
        if e.bbl and e.date in window and e.line_key.startswith("MEK#")]
    assert left == []


def test_a_fill_does_not_write_into_planned_downtime(db):
    """The optimizer creates no variable on an outage day, so a charge typed
    there is invisible to the solve and survives into the result. The referee
    catches that now; there is no reason to keep making it upstream."""
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    scenario, live, dates = _fill_scenario(db, "fill-downtime")
    line = sorted(k for k in live if k.startswith("ROSE#"))[0]
    window = dates[:10]
    svc.add_downtime(db, scenario.id, "ROSE", window[2], window[4],
                     "test outage", "test")
    outage = svc.downtime_days(db, scenario.id).get("ROSE", set())
    expected = sum(1 for d in window if d in outage)
    assert expected >= 3

    svc.fill_schedule(db, scenario.id, "ROSE", 0.0, live_keys=live, actor="test")
    r = svc.fill_schedule(db, scenario.id, line, 2500.0, start=window[0],
                          end=window[-1], live_keys=live, actor="test")
    assert r["skipped_downtime"] == expected

    got = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == line)}
    for d in window:
        if d in outage:
            assert not got.get(d), "{} charged while ROSE is down".format(d)
        else:
            assert got.get(d) == 2500.0


def test_filling_crude_sets_the_rate_and_the_mode(db):
    from invplanner.db import service as svc
    from invplanner.db.models import CrudeDay

    scenario, _, dates = _fill_scenario(db, "fill-crude")
    window = dates[:15]
    r = svc.fill_schedule(db, scenario.id, "CRUDE", 9500.0, start=window[0],
                          end=window[-1], mode="R", actor="test")
    assert r["cells"] == len(window)

    got = {c.date: c for c in db.query(CrudeDay).filter(
        CrudeDay.scenario_id == scenario.id)}
    for d in window:
        assert got[d].bbl == 9500.0
        assert got[d].mode == "R"


def test_only_empty_days_leaves_what_is_already_scheduled(db):
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    scenario, live, dates = _fill_scenario(db, "fill-empty-only")
    line = sorted(k for k in live if k.startswith("PLATFORMER#"))[0]
    down = svc.downtime_days(db, scenario.id).get("PLATFORMER", set())
    window = [d for d in dates[20:32] if d not in down]
    assert len(window) > 4

    svc.fill_schedule(db, scenario.id, "PLATFORMER", 0.0, live_keys=live,
                      actor="test")
    _set_cell(db, scenario.id, line, window[3], 777.0)

    svc.fill_schedule(db, scenario.id, line, 4000.0, start=window[0],
                      end=window[-1], overwrite=False, live_keys=live,
                      actor="test")
    got = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == line)}
    assert got[window[3]] == 777.0, "overwrote a day that was already scheduled"
    assert got[window[0]] == 4000.0


def test_the_fill_endpoint_rejects_a_retired_line(db):
    """The same guard the cell edit has. A bookmarked grid can still name a line
    the reference has since retired, and a fill would write hundreds of cells the
    optimizer then drops without saying so."""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db.session import get_session

    scenario, _, _ = _fill_scenario(db, "fill-retired")
    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    r = client.post("/api/scenarios/{}/schedule/fill".format(scenario.id),
                    json={"target": "MEK#70", "bbl": 1000.0})
    assert r.status_code == 400
    assert "retired" in r.json()["detail"]


def test_backfilling_gaps_leaves_the_schedule_it_fills_around(db):
    """`overwrite=False` plus the sibling clear is the rolling-plan case.

    A copied schedule runs out before the window does, and the fix is to write a
    rate onto the days that have none. If the sibling clear followed the date
    *range* rather than the days actually written, that backfill would delete
    every other line of the unit on the days the copy did cover - emptying the
    schedule it was called to complete.
    """
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    scenario, live, dates = _fill_scenario(db, "fill-backfill")
    hydro = sorted(k for k in live if k.startswith("HYDRO#"))
    assert len(hydro) > 2
    target, other = hydro[0], hydro[1]
    down = svc.downtime_days(db, scenario.id).get("HYDRO", set())
    window = [d for d in dates[:16] if d not in down]
    assert len(window) > 6

    svc.fill_schedule(db, scenario.id, "HYDRO", 0.0, live_keys=live, actor="test")
    # A covered stretch on one line, then nothing - the shape a copied schedule
    # has once the plan has rolled past the end of the source data.
    covered = window[:3]
    for d in covered:
        _set_cell(db, scenario.id, other, d, 4500.0)

    r = svc.fill_schedule(db, scenario.id, target, 2000.0, start=window[0],
                          end=window[-1], overwrite=False, live_keys=live,
                          actor="test")

    kept = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == other)}
    for d in covered:
        assert kept.get(d) == 4500.0, \
            "{} was cleared on a day the backfill skipped".format(other)
    assert r["siblings_cleared"] == 0

    got = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == target)}
    for d in window[3:]:
        assert got.get(d) == 2000.0, "gap not filled on {}".format(d)
    # The covered days are days the unit is already running, so they are not
    # gaps - "empty" is a property of the unit, not of one of its lines.
    for d in covered:
        assert not got.get(d), "wrote into a day the unit was already running"


def test_overwriting_still_clears_the_siblings_on_every_day_written(db):
    """The guard above must not have turned the sibling clear off. A plain fill
    still owns the whole range on a unit that runs one feed at a time."""
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    scenario, live, dates = _fill_scenario(db, "fill-overwrite-siblings")
    mek = sorted(k for k in live if k.startswith("MEK#"))
    assert len(mek) > 1, "this test needs a one-feed unit with more than one line"
    target, other = mek[0], mek[1]
    down = svc.downtime_days(db, scenario.id).get("MEK", set())
    window = [d for d in dates[:16] if d not in down]

    svc.fill_schedule(db, scenario.id, "MEK", 0.0, live_keys=live, actor="test")
    for d in window[:3]:
        _set_cell(db, scenario.id, other, d, 4500.0)

    r = svc.fill_schedule(db, scenario.id, target, 2000.0, start=window[0],
                          end=window[-1], live_keys=live, actor="test")
    assert r["siblings_cleared"] == 3

    kept = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
        ScheduleEntry.scenario_id == scenario.id,
        ScheduleEntry.line_key == other)}
    for d in window[:3]:
        assert not kept.get(d)


def test_a_fill_leaves_the_other_lines_where_several_run_on_one_day(db):
    """Clearing the siblings is right on MEK, EXTRACT and ROSE and wrong
    elsewhere. In the plans on 2026-09-13 the Platformer had two or more lines on
    3,573 of 5,968 charged days, HYDRO on 230 of 2,810 and TRANSFER_DIESEL on
    288 of 2,246. Filling one Platformer line cleared the others, so laying the
    Platformer down a line at a time wiped each line as the next was filled."""
    from invplanner import model_config as cfg
    from invplanner.db import service as svc
    from invplanner.db.models import ScheduleEntry

    assert cfg.ONE_FEED_UNITS == ["MEK", "EXTRACT", "ROSE"]
    scenario, live, dates = _fill_scenario(db, "fill-parallel-lines")
    for unit in ("PLATFORMER", "HYDRO", "TRANSFER_DIESEL"):
        assert not cfg.runs_one_feed(unit)
        lines = sorted(k for k in live if k.startswith(unit + "#"))
        assert len(lines) > 1, "{} needs more than one line".format(unit)
        target, other = lines[0], lines[1]
        down = svc.downtime_days(db, scenario.id).get(unit, set())
        window = [d for d in dates[:12] if d not in down]
        svc.fill_schedule(db, scenario.id, unit, 0.0, live_keys=live, actor="test")
        for d in window[:3]:
            _set_cell(db, scenario.id, other, d, 1500.0)

        r = svc.fill_schedule(db, scenario.id, target, 2000.0, start=window[0],
                              end=window[-1], live_keys=live, actor="test")
        assert r["siblings_cleared"] == 0, unit
        kept = {e.date: e.bbl for e in db.query(ScheduleEntry).filter(
            ScheduleEntry.scenario_id == scenario.id,
            ScheduleEntry.line_key == other)}
        for d in window[:3]:
            assert kept.get(d) == 1500.0, "{} lost {} on {}".format(unit, other, d)


# ------------------------------------------------------------ optimizer runs
def test_a_second_run_is_refused_while_one_is_solving(db):
    """Two v2 runs on the same plan were started 14 s apart, because the button
    stayed clickable through a 15-minute solve. The server refuses the second.
    A row left `running` by a server stopped mid-solve must not refuse runs for
    ever, and reads as interrupted."""
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import optimizer_service as optsvc
    from invplanner.db import service as svc
    from invplanner.db.models import OptimizerRun
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    base = svc.create_scenario(db, "run guard", dt.date(2026, 7, 23), 60, "test")
    params = dict(optsvc.params_dict(optsvc.active_params(db)),
                  time_limit_seconds=900)
    live = OptimizerRun(base_scenario_id=base.id, created_by="test",
                        params=params, status="running")
    db.add(live)
    db.commit()
    try:
        r = client.post("/api/optimizer/run", json={"scenario_id": base.id})
        assert r.status_code == 409, r.text
        assert "still solving" in r.json()["detail"]

        listed = {x["id"]: x for x in client.get("/api/optimizer/runs").json()}
        assert listed[live.id]["in_progress"] is True
        assert listed[live.id]["time_limit_seconds"] == 900
        # the run carries the settings it was given, labelled as on screen
        labels = {s["label"] for s in listed[live.id]["settings"]}
        assert "Units must run daily" in labels

        # a run that should have ended hours ago died with the server
        live.created_at = dt.datetime.utcnow() - dt.timedelta(hours=3)
        db.commit()
        assert optsvc.running_run(db) is None
        listed = {x["id"]: x for x in client.get("/api/optimizer/runs").json()}
        assert listed[live.id]["status"] == "interrupted"
        assert listed[live.id]["in_progress"] is False
    finally:
        live.status = "error"   # never leave a live-looking row for other tests
        db.commit()
        app.dependency_overrides.clear()


def test_a_schedule_the_simulator_rejected_is_not_exported(db):
    """An unverified run keeps its scenario so the failure can be inspected, but
    the schedule must not reach the workbook. The export refuses it, and the
    scenario list and the schedule's pair both say the run was rejected, which
    is what the picker and the compare bar mark."""
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import optimizer_service as optsvc
    from invplanner.db import service as svc
    from invplanner.db.models import OptimizerRun
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    base = svc.create_scenario(db, "rejected export", dt.date(2026, 7, 23), 60,
                               "test")

    class _Result:                       # what a solver hands back
        charge = {}
        transfer_bbl = {}

    result_id = optsvc._write_result_scenario(db, base, _Result(), "test", 0)
    run = OptimizerRun(base_scenario_id=base.id, result_scenario_id=result_id,
                       created_by="test", params={}, status="unverified",
                       message="The schedule asks units to charge feed the tanks "
                               "cannot supply.")
    db.add(run)
    db.commit()
    try:
        r = client.get("/api/scenarios/{}/schedule.xlsx?days=14".format(result_id))
        assert r.status_code == 409, r.text
        assert "rejected" in r.json()["detail"]
        # the plan it came from still exports
        assert client.get("/api/scenarios/{}/schedule.xlsx?days=14"
                          .format(base.id)).status_code == 200

        listed = {s["id"]: s for s in client.get("/api/scenarios").json()}
        assert listed[result_id]["run_status"] == "unverified"
        assert listed[base.id]["run_status"] is None
        pair = client.get("/api/scenarios/{}".format(result_id)).json()["pair"]
        assert pair["run_status"] == "unverified"
    finally:
        app.dependency_overrides.clear()


def test_results_are_labelled_from_their_run_and_their_current_plan(db):
    """Result names nested - "Optimizer run 21 (from Optimizer run 20 (from
    ...))" - until the picker pushed New scenario off the screen. A result is
    labelled from its run and the Current Plan at the root of the chain, new
    results are stored under that short name, and the schedule's pair speaks
    in Current Plan / Optimized Result."""
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import optimizer_service as optsvc
    from invplanner.db import service as svc
    from invplanner.db.models import OptimizerRun, Scenario
    from invplanner.db.session import get_session

    class _Result:                       # what a solver hands back
        charge = {}
        transfer_bbl = {}

    plan = svc.create_scenario(db, "Plan labels", dt.date(2026, 7, 23), 60, "test")
    first_id = optsvc._write_result_scenario(db, plan, _Result(), "test", 0,
                                             model_version="greedy")
    run1 = OptimizerRun(base_scenario_id=plan.id, result_scenario_id=first_id,
                        created_by="test", params={"model_version": "greedy"},
                        status="not_proved_optimal")
    db.add(run1)
    db.commit()
    # refined from the result, as "Refine with v2" will
    first = db.query(Scenario).get(first_id)
    refined_id = optsvc._write_result_scenario(db, first, _Result(), "test", 0,
                                               model_version="v2")
    run2 = OptimizerRun(base_scenario_id=first_id, result_scenario_id=refined_id,
                        created_by="test", params={"model_version": "v2"},
                        status="not_proved_optimal")
    db.add(run2)
    db.commit()

    labels = optsvc.scenario_labels(db)
    assert labels[plan.id]["kind"] == "current_plan"
    assert labels[plan.id]["label"] == "Plan labels"
    assert labels[first_id]["kind"] == "optimized_result"
    assert labels[first_id]["label"] == \
        "Run {} \u00b7 greedy \u00b7 from Plan labels".format(run1.id)
    # two deep, and still named after the plan rather than nested
    assert labels[refined_id]["label"] == \
        "Run {} \u00b7 v2 \u00b7 from Plan labels".format(run2.id)
    assert labels[refined_id]["plan_id"] == plan.id
    stored = db.query(Scenario).get(refined_id).name
    assert "(from" not in stored and stored.endswith("from Plan labels")

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    try:
        listed = {s["id"]: s for s in client.get("/api/scenarios").json()}
        assert listed[plan.id]["kind"] == "current_plan"
        assert listed[first_id]["kind"] == "optimized_result"

        on_result = client.get("/api/scenarios/{}".format(first_id)).json()
        assert on_result["plan_label"] == "Plan labels"
        assert on_result["pair"]["role"] == "optimized_result"
        assert on_result["pair"]["other_role"] == "current_plan"
        assert on_result["pair"]["other_label"] == "Plan labels"

        on_plan = client.get("/api/scenarios/{}".format(plan.id)).json()
        assert on_plan["pair"]["role"] == "current_plan"
        assert on_plan["pair"]["other_role"] == "optimized_result"
    finally:
        app.dependency_overrides.clear()


def test_a_plan_without_a_name_is_refused(db):
    """The form will not submit a blank name, and neither will the API: a
    blank entry in the picker is a plan nobody can find again."""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    try:
        before = len(client.get("/api/scenarios").json())
        r = client.post("/api/scenarios",
                        json={"name": "   ", "as_of": "2026-09-13"})
        assert r.status_code == 400, r.text
        assert len(client.get("/api/scenarios").json()) == before
    finally:
        app.dependency_overrides.clear()


def test_the_page_and_its_files_are_revalidated():
    """A browser kept an old index.html beside a new app.js, and the app stopped
    loading. The page and the static files tell the browser to check first."""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app

    client = TestClient(app)
    for path in ("/", "/static/app.js", "/static/styles.css"):
        r = client.get(path)
        assert r.status_code == 200, path
        assert r.headers.get("cache-control") == "no-cache", path
    # the API is left alone
    assert "cache-control" not in client.get("/api/health").headers


def test_undo_puts_a_removed_downtime_window_back_as_it_was(db):
    """Removing downtime is one click with Undo, and Undo adds the window again.
    A detected window - inferred from the workbook, never confirmed - has to come
    back detected, not quietly promoted to confirmed."""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import service as svc
    from invplanner.db.models import Scenario as ScenarioModel
    from invplanner.db.session import get_session

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    sid = db.query(ScenarioModel).order_by(ScenarioModel.id).first().id
    url = "/api/scenarios/{}/downtime".format(sid)
    try:
        r = client.post(url, json={"unit": "ROSE", "start": "2026-10-01",
                                   "end": "2026-10-03", "reason": "undo check",
                                   "status": "detected"})
        assert r.status_code == 200, r.text
        wins = {w["id"]: w for w in svc.list_downtime(db, sid)}
        assert wins[r.json()["id"]]["status"] == "detected"

        bad = client.post(url, json={"unit": "ROSE", "start": "2026-10-01",
                                     "end": "2026-10-03", "status": "maybe"})
        assert bad.status_code == 400

        # a window added by hand, with no status sent, is still confirmed
        by_hand = client.post(url, json={"unit": "ROSE", "start": "2026-10-05",
                                         "end": "2026-10-06"})
        wins = {w["id"]: w for w in svc.list_downtime(db, sid)}
        assert wins[by_hand.json()["id"]]["status"] == "confirmed"
    finally:
        for w in svc.list_downtime(db, sid):
            if w["unit"] == "ROSE" and w["start"] in ("2026-10-01", "2026-10-05"):
                client.delete("{}/{}".format(url, w["id"]))
        app.dependency_overrides.clear()


# ------------------------------------------------------------------ deleting
def _plan_with_a_refined_result(db, name):
    """A Current Plan, a greedy result from it, and a v2 result refined from that."""
    import datetime as dt

    from invplanner.db import optimizer_service as optsvc
    from invplanner.db import service as svc
    from invplanner.db.models import OptimizerRun, Scenario

    class _Result:                       # what a solver hands back
        charge = {}
        transfer_bbl = {}

    plan = svc.create_scenario(db, name, dt.date(2026, 7, 23), 60, "test")
    first_id = optsvc._write_result_scenario(db, plan, _Result(), "test", 0,
                                             model_version="greedy")
    run1 = OptimizerRun(base_scenario_id=plan.id, result_scenario_id=first_id,
                        created_by="test", params={"model_version": "greedy"},
                        status="not_proved_optimal", kpis={})
    db.add(run1)
    db.commit()
    refined_id = optsvc._write_result_scenario(
        db, db.query(Scenario).get(first_id), _Result(), "test", 0,
        model_version="v2")
    run2 = OptimizerRun(base_scenario_id=first_id, result_scenario_id=refined_id,
                        created_by="test", params={"model_version": "v2"},
                        status="not_proved_optimal", kpis={})
    db.add(run2)
    db.commit()
    return plan.id, first_id, refined_id, run1.id, run2.id


def test_deleting_a_current_plan_takes_its_results_and_run_cards(db):
    """A Current Plan goes with every Optimized Result made from it, however
    many refinements deep, every run card started from any of them, and their
    grids and downtime. Another plan is untouched. (Convention 13.)"""
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import service as svc
    from invplanner.db.models import (CrudeDay, OptimizerRun, PlannedDowntime,
                                      Scenario, ScheduleEntry)
    from invplanner.db.session import get_session

    plan, first, refined, run1, run2 = _plan_with_a_refined_result(db, "Plan to delete")
    failed = OptimizerRun(base_scenario_id=plan, created_by="test", params={},
                          status="infeasible")
    db.add(failed)
    db.commit()
    failed_id = failed.id
    other = svc.create_scenario(db, "Plan that stays", dt.date(2026, 7, 23), 60,
                                "test").id
    gone = [plan, first, refined]
    assert db.query(ScheduleEntry).filter(ScheduleEntry.scenario_id.in_(gone)).count()

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    try:
        preview = client.get("/api/scenarios/{}/deletion".format(plan)).json()
        assert preview["kind"] == "current_plan"
        assert set(preview["result_ids"]) == {first, refined}
        assert set(preview["runs"]) == {run1, run2, failed_id}
        assert preview["kept"] == [] and preview["busy"] == []

        r = client.delete("/api/scenarios/{}".format(plan))
        assert r.status_code == 200, r.text
        for sid in gone:
            assert db.query(Scenario).get(sid) is None
        assert db.query(OptimizerRun).filter(
            OptimizerRun.id.in_([run1, run2, failed_id])).count() == 0
        for model in (ScheduleEntry, CrudeDay, PlannedDowntime):
            assert db.query(model).filter(model.scenario_id.in_(gone)).count() == 0
        assert db.query(Scenario).get(other) is not None
        assert client.get("/api/scenarios/{}".format(plan)).status_code == 404
    finally:
        app.dependency_overrides.clear()


def test_deleting_a_result_keeps_the_results_made_from_it(db):
    """An Optimized Result goes with its own run card. A result refined from it
    stays, still names its Current Plan, keeps a valid reference, and its card
    still says which result it really started from. (Convention 13.)"""
    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import optimizer_service as optsvc
    from invplanner.db.models import OptimizerRun, Scenario
    from invplanner.db.session import get_session

    plan, first, refined, run1, run2 = _plan_with_a_refined_result(db, "Plan keeps")
    labels = optsvc.scenario_labels(db)

    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    try:
        preview = client.get("/api/scenarios/{}/deletion".format(first)).json()
        assert preview["kind"] == "optimized_result"
        assert preview["result_ids"] == [first] and preview["runs"] == [run1]
        assert preview["kept"] == [labels[refined]["label"]]
        assert preview["plan_label"] == "Plan keeps"

        assert client.delete("/api/scenarios/{}".format(first)).status_code == 200
        assert db.query(Scenario).get(first) is None
        assert db.query(OptimizerRun).get(run1) is None
        assert db.query(Scenario).get(plan) is not None
        assert db.query(Scenario).get(refined) is not None

        kept = db.query(OptimizerRun).get(run2)
        assert kept.base_scenario_id == plan
        assert kept.kpis["started_from_deleted"]["label"] == labels[first]["label"]
        assert optsvc.scenario_labels(db)[refined]["label"].endswith("from Plan keeps")
    finally:
        app.dependency_overrides.clear()


def test_a_plan_is_not_deleted_while_a_run_solves_on_it(db):
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import service as svc
    from invplanner.db.models import OptimizerRun, Scenario
    from invplanner.db.session import get_session

    plan = svc.create_scenario(db, "Plan still solving", dt.date(2026, 7, 23), 60,
                               "test")
    live = OptimizerRun(base_scenario_id=plan.id, created_by="test",
                        params={"time_limit_seconds": 900}, status="running")
    db.add(live)
    db.commit()
    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    try:
        r = client.delete("/api/scenarios/{}".format(plan.id))
        assert r.status_code == 409, r.text
        assert db.query(Scenario).get(plan.id) is not None
    finally:
        live.status = "error"   # never leave a live-looking row for other tests
        db.commit()
        app.dependency_overrides.clear()


def test_refine_with_v2_runs_v2_on_a_verified_greedy_result_only(db, monkeypatch):
    """Refine with v2 is the one way to solve an Optimized Result again
    (convention 12): v2, on a verified greedy run's result. An unverified greedy
    run and a non-greedy run are refused, and a card knows what was already
    refined from its result. The solve itself is stubbed out."""
    import datetime as dt

    from fastapi.testclient import TestClient

    from invplanner.api.main import app
    from invplanner.db import optimizer_service as optsvc
    from invplanner.db import service as svc
    from invplanner.db.models import OptimizerParams, OptimizerRun
    from invplanner.db.session import get_session

    class _Result:                       # what a solver hands back
        charge = {}
        transfer_bbl = {}

    plan = svc.create_scenario(db, "Plan to refine", dt.date(2026, 7, 23), 60, "test")
    result_id = optsvc._write_result_scenario(db, plan, _Result(), "test", 0,
                                              model_version="greedy")
    make = lambda model, status: OptimizerRun(  # noqa: E731
        base_scenario_id=plan.id, result_scenario_id=result_id, created_by="test",
        params={"model_version": model}, status=status, kpis={})
    greedy, rejected, solved = (make("greedy", "not_proved_optimal"),
                                make("greedy", "unverified"),
                                make("v2", "not_proved_optimal"))
    child = OptimizerRun(base_scenario_id=result_id, created_by="test",
                         params={"model_version": "v2"}, status="error")
    db.add_all([greedy, rejected, solved, child])
    db.commit()

    calls = []
    monkeypatch.setattr(optsvc, "run", lambda db_, base_id, actor="planner",
                        model_version=None: calls.append((base_id, model_version)))
    monkeypatch.setattr(OptimizerParams, "missing", lambda self: [])
    app.dependency_overrides[get_session] = lambda: db
    client = TestClient(app)
    try:
        listed = {x["id"]: x for x in client.get("/api/optimizer/runs").json()}
        assert listed[greedy.id]["refinable"] is True
        assert listed[rejected.id]["refinable"] is False
        assert listed[solved.id]["refinable"] is False
        assert listed[greedy.id]["refined_by"] == [child.id]

        url = "/api/optimizer/runs/{}/refine"
        assert client.post(url.format(rejected.id)).status_code == 409
        assert client.post(url.format(solved.id)).status_code == 409
        assert client.post(url.format(999999)).status_code == 404
        assert calls == []

        r = client.post(url.format(greedy.id))
        assert r.status_code == 200, r.text
        assert calls == [(result_id, "v2")]
    finally:
        app.dependency_overrides.clear()
